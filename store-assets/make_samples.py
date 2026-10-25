"""Builds the clean English sample documents shown in the store screenshots.

Run with a Python that has python-docx, openpyxl, python-pptx and reportlab:
    python make_samples.py samples/
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt as PPt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

BLUE = "1F6FEB"
out = Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
out.mkdir(parents=True, exist_ok=True)


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)


def word():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    t = d.add_heading("Project Proposal: Green Office 2027", 0)
    for r in t.runs:
        r.font.color.rgb = RGBColor.from_string(BLUE)
    p = d.add_paragraph("Prepared for the Board · October 2026")
    p.runs[0].italic = True
    d.add_heading("1. Summary", 1)
    d.add_paragraph(
        "Green Office 2027 cuts our paper use by 80% and our energy bill by a quarter. "
        "Every contract, invoice and report moves to digital documents that the whole team "
        "can open, edit and sign from their phone."
    )
    d.add_heading("2. Goals", 1)
    for g in ["Go paperless in all three offices by June 2027",
              "Sign every client contract digitally",
              "Save $48,000 a year on printing and storage"]:
        d.add_paragraph(g, style="List Bullet")
    d.add_heading("3. Timeline", 1)
    rows = [("Phase", "Months", "Owner"), ("Plan", "Jan – Feb", "Anna Lee"),
            ("Pilot", "Mar – Apr", "David Kim"), ("Roll-out", "May – Jun", "Maria Lopez")]
    tb = d.add_table(rows=len(rows), cols=3)
    tb.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tb.cell(i, j)
            c.text = v
            if i == 0:
                shade(c, BLUE)
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                c.paragraphs[0].runs[0].bold = True
            elif i % 2 == 0:
                shade(c, "EAF1FD")
    d.add_heading("4. Budget", 1)
    d.add_paragraph("The total budget is $120,000, covered by the savings within 30 months.")
    q = d.add_paragraph("“The best time to go paperless was yesterday. The next best time is today.”")
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.runs[0].italic = True
    d.save(out / "Project Proposal.docx")


def excel():
    # Large type and wide columns so the sheet fills a phone screen at the viewer's default zoom.
    wb = Workbook()
    ws = wb.active
    ws.title = "2026 Sales"
    big = 20
    ws["A1"] = "2026 Sales Report"
    ws["A1"].font = Font(size=30, bold=True, color="128A3C")
    ws.row_dimensions[1].height = 48
    head = ["Month", "Revenue", "Costs", "Profit"]
    months = ["January", "February", "March", "April", "May", "June", "July",
              "August", "September", "October", "November", "December"]
    revenue = [42800, 45100, 51900, 49600, 56300, 61200, 58700, 63400, 67100, 70800, 74500, 82900]
    costs = [28100, 29400, 31800, 30900, 33700, 35200, 34600, 36900, 38100, 39400, 41000, 44800]
    thin = Side(style="thin", color="D0D7E2")
    for j, h in enumerate(head, 1):
        c = ws.cell(3, j, h)
        c.font = Font(size=big, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="128A3C")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 36
    for i, (m, r, k) in enumerate(zip(months, revenue, costs), 4):
        # Plain numbers: openpyxl writes no cached formula results for the viewer to show.
        for j, v in enumerate([m, r, k, r - k], 1):
            c = ws.cell(i, j, v)
            c.font = Font(size=big, color="1E7B34" if j == 4 else "222222")
            if j > 1:
                c.number_format = "$#,##0"
            if i % 2:
                c.fill = PatternFill("solid", fgColor="E8F5EC")
        ws.row_dimensions[i].height = 34
    last = 3 + len(months)
    tot = last + 1
    ws.cell(tot, 1, "Total")
    for j, v in enumerate([sum(revenue), sum(costs), sum(revenue) - sum(costs)], 2):
        ws.cell(tot, j, v).number_format = "$#,##0"
    for j in range(1, 5):
        c = ws.cell(tot, j)
        c.font = Font(size=big, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0E6B2E")
    ws.row_dimensions[tot].height = 36
    for row in ws.iter_rows(min_row=3, max_row=tot, max_col=4):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.column_dimensions["A"].width = 26
    for col in "BCD":
        ws.column_dimensions[col].width = 22
    wb.save(out / "Sales Report 2026.xlsx")


def slides():
    pr = Presentation()
    pr.slide_width, pr.slide_height = Inches(13.333), Inches(7.5)
    colors = ["C9520A", "1F6FEB", "128A3C", "7A3FD1"]
    content = [("Product Launch 2027", "Folio Notes · Go-to-market plan"),
               ("Why now?", "Remote teams sign 3× more documents on mobile than in 2023"),
               ("Our plan", "Beta in March · Launch in May · 1M users by December"),
               ("Thank you!", "Questions? hello@example.com")]
    for k, (title, sub) in enumerate(content):
        s = pr.slides.add_slide(pr.slide_layouts[6])
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, pr.slide_width, pr.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = PRGB.from_string(colors[k])
        bg.line.fill.background()
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1.5), Inches(5.5), Inches(5.5))
        dot.fill.solid()
        # A lighter tint, not transparency: the viewer draws transparent fills opaque.
        base = [int(colors[k][i:i + 2], 16) for i in (0, 2, 4)]
        dot.fill.fore_color.rgb = PRGB(*[c + (255 - c) * 18 // 100 for c in base])
        dot.line.fill.background()
        for text, top, size, bold in [(title, 2.6, 54, True), (sub, 3.9, 24, False)]:
            tb = s.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11), Inches(1.2))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = text
            r.font.size = PPt(size)
            r.font.bold = bold
            r.font.color.rgb = PRGB(255, 255, 255)
    pr.save(out / "Product Launch.pptx")


def pdf():
    c = canvas.Canvas(str(out / "Rental Agreement.pdf"), pagesize=A4)
    w, h = A4
    c.setFillColor(HexColor("#1F6FEB"))
    c.rect(0, h - 90, w, 90, fill=1, stroke=0)
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Helvetica-Bold", 24)
    c.drawString(50, h - 58, "Residential Rental Agreement")
    c.setFillColor(HexColor("#222222"))
    y = h - 130
    paras = [
        ("1. Parties", "This agreement is made on October 25, 2026 between Sunrise Homes LLC (the Landlord) "
                       "and Emma Johnson (the Tenant) for the apartment at 24 Maple Street, Unit 5B."),
        ("2. Term", "The lease starts on November 1, 2026 and ends on October 31, 2027. "
                    "The Tenant may renew for another twelve months with 60 days written notice."),
        ("3. Rent", "Monthly rent is $1,450, due on the first day of each month. "
                    "A security deposit of $1,450 is paid on signing and returned within 30 days of move-out."),
        ("4. Utilities", "The Tenant pays electricity and internet. Water, heating and trash are included."),
        ("5. Pets", "One cat or small dog is allowed with a one-time pet fee of $250."),
        ("6. Maintenance", "The Landlord repairs appliances and plumbing within 72 hours of notice. "
                           "The Tenant keeps the unit clean and reports any damage promptly."),
    ]
    for head, body in paras:
        c.setFont("Helvetica-Bold", 13)
        c.drawString(50, y, head)
        y -= 20
        c.setFont("Helvetica", 11)
        line = ""
        for word_ in body.split():
            if c.stringWidth(line + " " + word_, "Helvetica", 11) > w - 100:
                c.drawString(50, y, line.strip())
                y -= 16
                line = ""
            line += " " + word_
        c.drawString(50, y, line.strip())
        y -= 30
    y -= 20
    c.setStrokeColor(HexColor("#888888"))
    for x, who in [(50, "Landlord"), (320, "Tenant")]:
        c.line(x, y, x + 220, y)
        c.setFont("Helvetica", 10)
        c.drawString(x, y - 14, f"{who} signature")
        c.line(x, y - 50, x + 220, y - 50)
        c.drawString(x, y - 64, "Date")
    c.showPage()
    c.save()


def text():
    (out / "Meeting Notes.txt").write_text(
        "Team meeting - Monday\n\n- Launch date moved to May 12\n- Anna sends the new budget by Friday\n"
        "- David books the venue\n- Next meeting: Thursday 10:00\n", encoding="utf-8")


word(); excel(); slides(); pdf(); text()
print("written to", out.resolve())
